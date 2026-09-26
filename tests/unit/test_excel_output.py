import io
import os
import sys

import openpyxl

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..')))

from model.compare_labels import (
    compare_labels, summarize, summary_header, compare_labels_by_region,
    build_region_summary_rows, COLOR_A_ONLY, COLOR_B_ONLY, COLOR_MISMATCH,
)
from model.excel_output import (
    create_compare_excel_output, create_region_compare_excel_output, region_sheet_names,
)


def _load(data):
    return openpyxl.load_workbook(io.BytesIO(data))


def _values(ws):
    return [[c.value for c in row] for row in ws.iter_rows()]


def test_sheet_names_and_columns():
    """シート名は Summary → All Labels diff、比較結果の列に機器符号候補を含む。"""
    df = compare_labels({'CB1': 1}, {'CB1': 1, '電源': 2})
    wb = _load(create_compare_excel_output(df, summarize(df, summary_header(1, 1, '全部', 1))))
    assert wb.sheetnames == ['Summary', 'All Labels diff']
    rows = _values(wb['All Labels diff'])
    assert rows[0] == ['機器符号候補', 'ラベル', '区分', 'A個数', 'B個数']
    assert rows[1] == ['Y', 'CB1', '両方', 1, 1]
    assert rows[2] == [None, '電源', 'B のみ', None, 2]


def test_summary_numbers_use_thousands_separator():
    """Summary の数値セルは数値型のまま #,##0 書式、文字列はそのまま。"""
    df = compare_labels({f'L{i}': 1 for i in range(1234)}, {})
    wb = _load(create_compare_excel_output(df, summarize(df, summary_header(1, 2, '全部', 2))))
    ws = wb['Summary']
    cells = {ws.cell(r, 1).value: ws.cell(r, 2) for r in range(2, ws.max_row + 1)}
    assert cells['A ユニークラベル数'].value == 1234
    assert cells['A ユニークラベル数'].number_format == '#,##0'
    assert cells['B 絞り込み条件'].value == '全部'


def test_row_colors():
    """行の塗り: Aのみ=青／Bのみ=緑／両方で個数不一致=黄／一致=無色。"""
    df = compare_labels({'A1': 1, 'M': 1, 'X': 1}, {'B1': 1, 'M': 1, 'X': 2})
    wb = _load(create_compare_excel_output(df, summarize(df, summary_header(1, 1, '全部', 1))))
    ws = wb['All Labels diff']
    fills = {ws.cell(r, 2).value: ws.cell(r, 2).fill.fgColor.rgb for r in range(2, ws.max_row + 1)}
    assert fills['A1'].endswith(COLOR_A_ONLY['bg_color'][1:])
    assert fills['B1'].endswith(COLOR_B_ONLY['bg_color'][1:])
    assert fills['X'].endswith(COLOR_MISMATCH['bg_color'][1:])
    assert fills['M'] in (None, '00000000')


def test_region_output_one_sheet_per_region():
    """指定領域での比較: 領域ごとのシート（シート名＝領域名、regions の順、領域名列なし）。
    該当ラベルが無い領域もヘッダーのみのシートとして出す。Summary は3列。"""
    regions = ['R2', 'SYSTEM I/F BOX', 'EMPTY']
    df, metrics = compare_labels_by_region(
        {'R2': {'A': 1, 'B': 1}, 'SYSTEM I/F BOX': {'CB1': 1}},
        {'R2': {'A': 1}}, regions)
    summary = build_region_summary_rows(metrics, summary_header(1, 1, '全部', 1))
    wb = _load(create_region_compare_excel_output(df, summary, regions))
    assert wb.sheetnames == ['Summary', 'R2', 'SYSTEM I_F BOX', 'EMPTY']
    assert _values(wb['Summary'])[0] == ['領域名', '項目', '値']
    header = ['機器符号候補', 'ラベル', '区分', 'A個数', 'B個数']
    r2 = _values(wb['R2'])
    assert r2[0] == header and [r[1] for r in r2[1:]] == ['A', 'B']
    assert [r[1] for r in _values(wb['SYSTEM I_F BOX'])[1:]] == ['CB1']
    assert _values(wb['EMPTY']) == [header]


def test_region_sheet_names_sanitize_truncate_and_dedupe():
    """Excelで使えない文字は '_'、31文字に切り詰め、Summary・他シートと重複すれば (2)…。"""
    names = region_sheet_names(
        ['SYSTEM I/F BOX', 'NXSAF ** (NX-ECC202)', 'summary', 'X:Y', 'X?Y', 'A' * 40])
    assert names['SYSTEM I/F BOX'] == 'SYSTEM I_F BOX'
    assert names['NXSAF ** (NX-ECC202)'] == 'NXSAF __ (NX-ECC202)'
    assert names['summary'] == 'summary(2)'
    assert (names['X:Y'], names['X?Y']) == ('X_Y', 'X_Y(2)')
    assert names['A' * 40] == 'A' * 31
    assert all(len(n) <= 31 for n in names.values())
